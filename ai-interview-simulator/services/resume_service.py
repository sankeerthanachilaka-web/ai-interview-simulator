from pathlib import Path
from uuid import uuid4
import re

from docx import Document
from pypdf import PdfReader
from werkzeug.utils import secure_filename

from config import Config
from models import Resume


class ResumeServiceError(Exception):
    pass


class ResumeService:

    # ==========================================================
    # ALLOWED FILES
    # ==========================================================

    def allowed(self, filename):

        if "." not in filename:
            return False

        extension = filename.rsplit(".", 1)[1].lower()

        return extension in {
            "pdf",
            "doc",
            "docx",
        }

    # ==========================================================
    # SAVE + EXTRACT
    # ==========================================================

    def save_and_extract(self, upload, user_id):

        filename = secure_filename(
            upload.filename or ""
        )

        if not filename:
            raise ResumeServiceError(
                "Please select a resume file."
            )

        if not self.allowed(filename):
            raise ResumeServiceError(
                "Only PDF, DOC, and DOCX resumes are supported."
            )

        ext = filename.rsplit(
            ".",
            1
        )[1].lower()

        safe_name = (
            f"{user_id}_"
            f"{uuid4().hex}."
            f"{ext}"
        )

        upload_folder = Path(
            Config.UPLOAD_FOLDER
        )

        upload_folder.mkdir(
            parents=True,
            exist_ok=True
        )

        destination = (
            upload_folder /
            safe_name
        )

        # ------------------------------------------------------
        # SAVE FILE
        # ------------------------------------------------------

        try:

            upload.save(
                str(destination)
            )

        except Exception as exc:

            raise ResumeServiceError(
                f"Could not save the resume: {exc}"
            )

        # ------------------------------------------------------
        # EXTRACT TEXT
        # ------------------------------------------------------

        try:

            if ext == "pdf":

                text = self._pdf(
                    destination
                )

            elif ext == "docx":

                text = self._docx(
                    destination
                )

            elif ext == "doc":

                raise ResumeServiceError(
                    "Old .DOC files cannot be reliably read "
                    "with the current extractor. Please save "
                    "your resume as .DOCX or .PDF and upload again."
                )

            else:

                raise ResumeServiceError(
                    "Unsupported resume format."
                )

        except ResumeServiceError:
            destination.unlink(
                missing_ok=True
            )
            raise

        except Exception as exc:

            destination.unlink(
                missing_ok=True
            )

            raise ResumeServiceError(
                "Could not read this resume. "
                "Please make sure the PDF/DOCX is valid."
            ) from exc

        # ------------------------------------------------------
        # CLEAN TEXT
        # ------------------------------------------------------

        text = self._clean_text(
            text
        )

        # ------------------------------------------------------
        # CHECK TEXT
        # ------------------------------------------------------

        if len(text) < 50:

            destination.unlink(
                missing_ok=True
            )

            raise ResumeServiceError(
                "No readable text was found in this resume. "
                "If your PDF is a scanned/image-only resume, "
                "please upload a text-based PDF or DOCX."
            )

        # ------------------------------------------------------
        # LIMIT SIZE
        # ------------------------------------------------------

        text = text[:50000]

        # ------------------------------------------------------
        # CREATE DATABASE OBJECT
        # ------------------------------------------------------

        return Resume(

            user_id=user_id,

            filename=filename,

            extracted_text=text,

        )

    # ==========================================================
    # PDF EXTRACTION
    # ==========================================================

    def _pdf(self, path):

        reader = PdfReader(
            str(path)
        )

        extracted_pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            try:

                page_text = (
                    page.extract_text(
                        extraction_mode="layout"
                    )
                    or ""
                )

            except TypeError:

                # Compatibility with older pypdf versions
                page_text = (
                    page.extract_text()
                    or ""
                )

            except Exception:

                page_text = ""

            if page_text.strip():

                extracted_pages.append(
                    f"\n--- PAGE {page_number} ---\n"
                    f"{page_text}"
                )

        return "\n".join(
            extracted_pages
        )

    # ==========================================================
    # DOCX EXTRACTION
    # ==========================================================

    def _docx(self, path):

        document = Document(
            str(path)
        )

        parts = []

        # ------------------------------------------------------
        # NORMAL PARAGRAPHS
        # ------------------------------------------------------

        for paragraph in document.paragraphs: 

            text = paragraph.text.strip()

            if text:

                parts.append(
                    text
                )

        # ------------------------------------------------------
        # TABLES
        # ------------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:

                        cells.append(
                            cell_text
                        )

                if cells:

                    parts.append(
                        " | ".join(cells)
                    )

        return "\n".join(
            parts
        )

    # ==========================================================
    # CLEAN TEXT
    # ==========================================================

    def _clean_text(self, text):

        if not text:
            return ""

        # Replace null characters
        text = text.replace(
            "\x00",
            " "
        )

        # Normalize whitespace
        text = re.sub(
            r"[ \t]+",
            " ",
            text
        )

        text = re.sub(
            r"\n\s*\n+",
            "\n\n",
            text
        )
 
        # Remove excessive repeated characters
        text = re.sub(
            r"\n{4,}",
            "\n\n",
            text
        )

        return text.strip()
